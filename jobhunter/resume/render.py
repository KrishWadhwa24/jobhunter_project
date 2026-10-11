"""ResumeDoc -> PDF (reportlab) and DOCX (python-docx). One column, standard headings, real text (no images)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

FONT = "Times-Roman"
FONT_B = "Times-Bold"
FONT_I = "Times-Italic"

SEC_ORDER = ["SUMMARY", "EXPERIENCE", "PROJECTS", "TECHNICAL SKILLS", "EDUCATION", "ACHIEVEMENTS"]


@dataclass
class Entry:
    heading: str                 # "Company — Title" or project name
    right: str = ""              # dates
    sub: str = ""                # location or tech stack
    links: List[Tuple[str, str]] = field(default_factory=list)
    bullets: List[str] = field(default_factory=list)


@dataclass
class ResumeDoc:
    name: str
    headline: str
    contact: List[str]                                   # plain pieces: location, phone, email
    links: List[Tuple[str, str]]                         # (label, url)
    summary: str
    experience: List[Entry]
    projects: List[Entry]
    skills: List[Tuple[str, List[str]]]
    education: Entry
    achievements: List[str]
    skills_title: str = "TECHNICAL SKILLS"
    experience_title: str = "EXPERIENCE"
    projects_title: str = "PROJECTS"
    achievements_title: str = "ACHIEVEMENTS"

    def to_text(self) -> str:
        out = [self.name, self.headline, " | ".join(self.contact + [l for l, _ in self.links]), "", "SUMMARY", self.summary, ""]
        if self.experience:
            out += [self.experience_title]
        for e in self.experience:
            out += [f"{e.heading} {e.right}"] + ([e.sub] if e.sub else []) + ["- " + b for b in e.bullets]
        if self.projects:
            out += ["", self.projects_title]
        for p in self.projects:
            out += [(f"{p.heading} | {p.sub}" if p.sub else p.heading)] + ["- " + b for b in p.bullets]
        out += ["", self.skills_title] + [f"{c}: {', '.join(items)}" for c, items in self.skills]
        out += ["", "EDUCATION", f"{self.education.heading} {self.education.right}", self.education.sub]
        if self.achievements:
            out += ["", self.achievements_title] + ["- " + a for a in self.achievements]
        return "\n".join(x for x in out if x is not None)


_PDF_MAP = {"–": "–", "—": "—", "‘": "'", "’": "'", "“": '"', "”": '"',
            "→": "->", "•": "-", "★": "*", "→": "->", " ": " ", " ": " "}


def pdf_safe(s: str) -> str:
    s = "".join(_PDF_MAP.get(c, c) for c in s)
    return s.encode("cp1252", "replace").decode("cp1252")


def _esc(s: str) -> str:
    return pdf_safe(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ----------------------------------------------------------------------------- PDF
def build_pdf(doc: ResumeDoc, scale: float = 1.0) -> Tuple[bytes, int]:
    """Returns (pdf_bytes, page_count)."""
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base = 9.4 * scale
    lead = base * 1.2
    st = {
        "name": ParagraphStyle("name", fontName=FONT_B, fontSize=16 * scale, leading=18 * scale, alignment=TA_CENTER,
                               spaceAfter=1),
        "headline": ParagraphStyle("hl", fontName=FONT_I, fontSize=base, leading=lead, alignment=TA_CENTER),
        "contact": ParagraphStyle("ct", fontName=FONT, fontSize=base - 0.2, leading=lead, alignment=TA_CENTER,
                                  spaceAfter=2),
        "h": ParagraphStyle("h", fontName=FONT_B, fontSize=base + 0.6, leading=lead + 0.6, spaceBefore=4.5, spaceAfter=0),
        "body": ParagraphStyle("b", fontName=FONT, fontSize=base, leading=lead, alignment=TA_JUSTIFY),
        "left": ParagraphStyle("l", fontName=FONT, fontSize=base, leading=lead, alignment=TA_LEFT),
        "sub": ParagraphStyle("s", fontName=FONT_I, fontSize=base, leading=lead, alignment=TA_LEFT),
        "bullet": ParagraphStyle("bu", fontName=FONT, fontSize=base, leading=lead, alignment=TA_JUSTIFY,
                                 leftIndent=11, bulletIndent=2, spaceBefore=0.4),
        "right": ParagraphStyle("r", fontName=FONT, fontSize=base, leading=lead, alignment=2),
    }
    buf = io.BytesIO()
    margin_x = 0.45 * inch
    pdf = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=margin_x, rightMargin=margin_x, topMargin=0.38 * inch,
                            bottomMargin=0.34 * inch, title=f"{doc.name} - Resume", author=doc.name)
    width = LETTER[0] - 2 * margin_x
    flow: List[Any] = []

    def link(label: str, url: str) -> str:
        return f'<a href="{_esc(url)}" color="#000000">{_esc(label)}</a>'

    flow.append(Paragraph(_esc(doc.name.upper()), st["name"]))
    if doc.headline:
        flow.append(Paragraph(_esc(doc.headline), st["headline"]))
    contact = [_esc(c) for c in doc.contact] + [link(l, u) for l, u in doc.links]
    flow.append(Paragraph(" | ".join(contact), st["contact"]))

    def heading(t: str) -> None:
        flow.append(Paragraph(t, st["h"]))
        flow.append(HRFlowable(width="100%", thickness=0.6, color="#000000", spaceBefore=0.5, spaceAfter=2))

    def two_col(left: str, right: str, left_style: str = "left") -> None:
        t = Table([[Paragraph(left, st[left_style]), Paragraph(_esc(right), st["right"])]],
                  colWidths=[width * 0.76, width * 0.24])
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        flow.append(t)

    heading("SUMMARY")
    flow.append(Paragraph(_esc(doc.summary), st["body"]))

    if doc.experience:
        heading(doc.experience_title)
    for e in doc.experience:
        two_col(f"<b>{_esc(e.heading)}</b>", e.right)
        if e.sub:
            flow.append(Paragraph(_esc(e.sub), st["sub"]))
        for b in e.bullets:
            flow.append(Paragraph(_esc(b), st["bullet"], bulletText="•"))

    if doc.projects:
        heading(doc.projects_title)
        for p in doc.projects:
            right = " | ".join(link(l, u) for l, u in p.links)
            t = Table([[Paragraph(f"<b>{_esc(p.heading)}</b> | <i>{_esc(p.sub)}</i>", st["left"]),
                        Paragraph(right, st["right"])]], colWidths=[width * 0.82, width * 0.18])
            t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(Spacer(1, 1.5))
            flow.append(t)
            for b in p.bullets:
                flow.append(Paragraph(_esc(b), st["bullet"], bulletText="•"))

    heading(doc.skills_title)
    for cat, items in doc.skills:
        flow.append(Paragraph(f"<b>{_esc(cat)}:</b> {_esc(', '.join(items))}", st["left"]))

    heading("EDUCATION")
    two_col(f"<b>{_esc(doc.education.heading)}</b>", doc.education.right)
    flow.append(Paragraph(_esc(doc.education.sub), st["left"]))

    if doc.achievements:
        heading(doc.achievements_title)
        for a in doc.achievements:
            flow.append(Paragraph(_esc(a), st["bullet"], bulletText="•"))

    pdf.build(flow)
    data = buf.getvalue()
    return data, _page_count(data)


def _page_count(data: bytes) -> int:
    try:
        from pypdf import PdfReader
        return len(PdfReader(io.BytesIO(data)).pages)
    except Exception:
        return len(re.findall(rb"/Type\s*/Page[^s]", data))


def pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    except Exception:
        return ""


# ----------------------------------------------------------------------------- DOCX
def build_docx(doc: ResumeDoc, scale: float = 1.0) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.45)
    sec.top_margin, sec.bottom_margin = Inches(0.38), Inches(0.34)
    base = 9.4 * scale
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(base)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    text_width = Inches(8.5 - 0.9)

    def para(align=None, space_before=0.0, space_after=0.0):
        p = d.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def run(p, text, bold=False, italic=False, size=None):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        return r

    def add_link(p, label: str, url: str) -> None:
        rid = d.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = label
        r.append(t)
        h.append(r)
        p._p.append(h)

    def rule(p) -> None:
        ppr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
            bt.set(qn("w:" + k), v)
        b.append(bt)
        ppr.append(b)

    def heading(t: str) -> None:
        p = para(space_before=4.5, space_after=2)
        run(p, t, bold=True, size=base + 0.6)
        rule(p)

    def tabbed(left_runs, right: str):
        p = para()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        for txt, b, i in left_runs:
            run(p, txt, bold=b, italic=i)
        run(p, "\t" + right)
        return p

    def bullet(text: str) -> None:
        p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0.4)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        run(p, "• " + text)

    p = para(WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    run(p, doc.name.upper(), bold=True, size=16 * scale)
    if doc.headline:
        p = para(WD_ALIGN_PARAGRAPH.CENTER)
        run(p, doc.headline, italic=True)
    p = para(WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    run(p, " | ".join(doc.contact) + (" | " if doc.links else ""))
    for i, (l, u) in enumerate(doc.links):
        add_link(p, l, u)
        if i < len(doc.links) - 1:
            run(p, " | ")

    heading("SUMMARY")
    p = para(WD_ALIGN_PARAGRAPH.JUSTIFY)
    run(p, doc.summary)

    if doc.experience:
        heading(doc.experience_title)
    for e in doc.experience:
        tabbed([(e.heading, True, False)], e.right)
        if e.sub:
            p = para()
            run(p, e.sub, italic=True)
        for b in e.bullets:
            bullet(b)

    if doc.projects:
        heading(doc.projects_title)
        for pr in doc.projects:
            p = para(space_before=1.5)
            p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            run(p, pr.heading, bold=True)
            run(p, " | ")
            run(p, pr.sub, italic=True)
            run(p, "\t")
            for i, (l, u) in enumerate(pr.links):
                add_link(p, l, u)
                if i < len(pr.links) - 1:
                    run(p, " | ")
            for b in pr.bullets:
                bullet(b)

    heading(doc.skills_title)
    for cat, items in doc.skills:
        p = para()
        run(p, cat + ": ", bold=True)
        run(p, ", ".join(items))

    heading("EDUCATION")
    tabbed([(doc.education.heading, True, False)], doc.education.right)
    p = para()
    run(p, doc.education.sub)

    if doc.achievements:
        heading(doc.achievements_title)
        for a in doc.achievements:
            bullet(a)

    d.core_properties.title = f"{doc.name} - Resume"
    d.core_properties.author = doc.name
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
